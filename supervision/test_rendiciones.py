from datetime import datetime, timezone
from io import BytesIO
from decimal import Decimal

from openpyxl import load_workbook
from pypdf import PdfReader
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from PIL import Image
from rest_framework.test import APIClient

from user.models import Cargo, Usuario
from .models import RendicionGasto


class RendicionGastoTests(TestCase):
    endpoint = "/api/supervision/rendiciones-gastos/"

    def setUp(self):
        cargo = Cargo.objects.create(nombre="Supervisor")
        self.supervisor = Usuario.objects.create_user(
            username="sup-gastos", nombres="Ana", apellidos="Pérez", rut="11111111-1",
            email="sup@example.test", cargo=cargo,
        )
        self.other = Usuario.objects.create_user(
            username="otro-gastos", nombres="Luis", apellidos="Soto", rut="22222222-2",
            email="otro@example.test", cargo=cargo,
        )
        self.admin = Usuario.objects.create_user(
            username="admin-gastos", rut="33333333-3", email="admin@example.test", is_staff=True,
        )
        self.client = APIClient()
        self.client.force_authenticate(self.supervisor)

    def image(self):
        buffer = BytesIO()
        Image.new("RGB", (20, 20), "white").save(buffer, format="PNG")
        return SimpleUploadedFile("comprobante.png", buffer.getvalue(), content_type="image/png")

    def expense(self, supervisor=None, **values):
        user = supervisor or self.supervisor
        return RendicionGasto.objects.create(
            supervisor=user, supervisor_nombre=str(user), storage_key="rendiciones-gastos/test.png",
            instalacion=values.get("instalacion", "Planta Norte"), persona=values.get("persona", "Pedro Soto"),
            monto=values.get("monto", 15000), motivo=values.get("motivo", "Traslado de guardia"),
            gasto_depositado=values.get("gasto_depositado", False), depositante=values.get("depositante", ""),
            monto_depositado=values.get("monto_depositado"),
        )

    @patch("supervision.rendiciones.upload_document")
    def test_deposits_persist_and_difference_is_computed_by_server(self, upload):
        for deposited, difference in [(10000, "5000"), (15000, "0"), (20000, "-5000")]:
            with self.subTest(deposited=deposited):
                response = self.client.post(self.endpoint, {
                    "imagen": self.image(), "instalacion": "Planta", "persona": "Pedro",
                    "monto": "15000", "motivo": "Traslado", "gasto_depositado": "true",
                    "depositante": "  Carlos Jefe  ", "monto_depositado": str(deposited), "diferencia": "999",
                }, format="multipart")
                self.assertEqual(response.status_code, 201, response.data)
                saved = RendicionGasto.objects.get(pk=response.data["id"])
                self.assertTrue(saved.gasto_depositado)
                self.assertEqual(saved.depositante, "Carlos Jefe")
                self.assertEqual(saved.monto_depositado, Decimal(deposited))
                detail = self.client.get(f"{self.endpoint}{saved.pk}/").data
                self.assertEqual(detail["diferencia"], difference)

    @patch("supervision.rendiciones.upload_document")
    def test_deposit_requires_name_and_positive_integer_amount_before_upload(self, upload):
        cases = [{"depositante": "  "}, {"depositante": None}, {"depositante": "x" * 221},
                 {"monto_depositado": None}, {"monto_depositado": "0"}, {"monto_depositado": "-1"},
                 {"monto_depositado": "1.5"}, {"monto_depositado": "abc"}, {"monto_depositado": "1000000000000"}]
        for changes in cases:
            payload = {"imagen": self.image(), "instalacion": "Planta", "persona": "Pedro", "monto": "15000",
                       "motivo": "Traslado", "gasto_depositado": "true", "depositante": "Carlos", "monto_depositado": "10000"}
            payload.update(changes)
            with self.subTest(changes=changes):
                response = self.client.post(self.endpoint, {k: v for k, v in payload.items() if v is not None}, format="multipart")
                self.assertEqual(response.status_code, 400, response.data)
        upload.assert_not_called()

    @patch("supervision.rendiciones.upload_document")
    def test_unchecked_deposit_clears_details_and_old_clients_remain_supported(self, upload):
        for deposit in [{}, {"gasto_depositado": "false", "depositante": "Carlos", "monto_depositado": "10000"}]:
            response = self.client.post(self.endpoint, {
                "imagen": self.image(), "instalacion": "Planta", "persona": "Pedro",
                "monto": "15000", "motivo": "Traslado", **deposit,
            }, format="multipart")
            self.assertEqual(response.status_code, 201, response.data)
            self.assertFalse(response.data["gasto_depositado"])
            self.assertEqual(response.data["depositante"], "")
            self.assertIsNone(response.data["monto_depositado"])
            self.assertIsNone(response.data["diferencia"])

    def test_reports_include_deposits_signed_differences_and_separate_balances(self):
        for amount in [10000, 15000, 20000]:
            self.expense(gasto_depositado=True, depositante="=Carlos <Jefe>", monto_depositado=amount)
        self.expense(monto=None)
        book = load_workbook(BytesIO(self.report().content))
        detail = book["Detalle"]
        self.assertEqual([detail.cell(i, 11).value for i in range(2, 6)], [5000, 0, -5000, None])
        self.assertEqual(detail["I2"].data_type, "s")
        self.assertEqual(detail["I2"].value, "=Carlos <Jefe>")
        self.assertEqual(detail["J2"].data_type, "n")
        self.assertEqual(detail["J6"].value, 45000)
        summary = dict(book["Resumen"].iter_rows(values_only=True))
        self.assertEqual(summary["Total depositado CLP"], 45000)
        self.assertEqual(summary["Por cubrir en gastos con depósito CLP"], 5000)
        self.assertEqual(summary["Sobrante en gastos con depósito CLP"], 5000)
        pdf = PdfReader(BytesIO(self.report("pdf").content))
        text = " ".join(page.extract_text() for page in pdf.pages)
        for expected in ["Depósitos y diferencias", "=Carlos <Jefe>", "45.000", "-5.000", "No calculada"]:
            self.assertIn(expected, text)

    @patch("supervision.rendiciones.upload_document")
    def test_create_records_authenticated_supervisor_and_server_time(self, upload):
        before = datetime.now(timezone.utc)
        response = self.client.post(self.endpoint, {
            "monto": "15000", "motivo": "  Traslado de guardia  ",
            "imagen": self.image(), "instalacion": "  Planta Norte  ", "persona": " Pedro Soto ",
            "supervisor": self.other.id, "supervisor_nombre": "Falso", "creada_en": "2000-01-01T00:00:00Z",
        }, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        expense = RendicionGasto.objects.get()
        self.assertEqual(expense.supervisor, self.supervisor)
        self.assertEqual(expense.supervisor_nombre, "Ana Pérez")
        self.assertGreaterEqual(expense.creada_en, before)
        self.assertEqual(expense.instalacion, "Planta Norte")
        self.assertEqual(expense.persona, "Pedro Soto")
        self.assertEqual(expense.monto, Decimal("15000"))
        self.assertEqual(expense.motivo, "Traslado de guardia")
        self.assertTrue(expense.storage_key.startswith("rendiciones-gastos/"))
        upload.assert_called_once()

    @patch("supervision.rendiciones.upload_document")
    def test_all_fields_required_and_invalid_images_rejected(self, upload):
        payloads = [
            {"instalacion": "Planta", "persona": "Pedro"},
            {"imagen": self.image(), "instalacion": "  ", "persona": "Pedro"},
            {"imagen": self.image(), "instalacion": "Planta", "persona": "  "},
            {"imagen": SimpleUploadedFile("fake.png", b"invalid", content_type="image/png"), "instalacion": "Planta", "persona": "Pedro"},
        ]
        for payload in payloads:
            payload.update(monto="15000", motivo="Traslado")
            with self.subTest(fields=list(payload)):
                self.assertEqual(self.client.post(self.endpoint, payload, format="multipart").status_code, 400)
        self.assertFalse(RendicionGasto.objects.exists())
        upload.assert_not_called()

    @patch("supervision.rendiciones.upload_document")
    def test_oversized_image_rejected(self, upload):
        image = self.image()
        oversized = SimpleUploadedFile("large.png", image.read() + b"\0" * (10 * 1024 * 1024), content_type="image/png")
        response = self.client.post(self.endpoint, {"imagen": oversized, "instalacion": "Planta", "persona": "Pedro", "monto": "15000", "motivo": "Traslado"}, format="multipart")
        self.assertEqual(response.status_code, 400)
        upload.assert_not_called()

    @patch("supervision.rendiciones.upload_document", side_effect=RuntimeError("Storage unavailable"))
    def test_storage_failure_does_not_create_expense(self, upload):
        with self.assertLogs("supervision.rendiciones", level="ERROR"):
            response = self.client.post(self.endpoint, {"imagen": self.image(), "instalacion": "Planta", "persona": "Pedro", "monto": "15000", "motivo": "Traslado"}, format="multipart")
        self.assertEqual(response.status_code, 500)
        self.assertFalse(RendicionGasto.objects.exists())

    @patch("supervision.rendiciones.upload_document")
    @patch("supervision.rendiciones.delete_document")
    @patch("supervision.rendiciones.RendicionSerializer.save", side_effect=RuntimeError("DB failure"))
    def test_database_failure_cleans_uploaded_image(self, save, delete, upload):
        with self.assertRaises(RuntimeError):
            self.client.post(self.endpoint, {"imagen": self.image(), "instalacion": "Planta", "persona": "Pedro", "monto": "15000", "motivo": "Traslado"}, format="multipart")
        delete.assert_called_once_with(upload.call_args.args[1])

    @patch("supervision.rendiciones.generate_signed_url", return_value="https://storage.test/comprobante")
    def test_supervisor_only_sees_own_expenses_and_images(self, sign):
        own = self.expense()
        other = self.expense(self.other)
        response = self.client.get(self.endpoint)
        self.assertEqual([row["id"] for row in response.data["results"]], [own.id])
        self.assertEqual(self.client.get(f"{self.endpoint}{other.id}/").status_code, 404)
        self.assertEqual(self.client.get(f"{self.endpoint}{other.id}/imagen/").status_code, 404)
        sign.assert_not_called()
        self.assertEqual(self.client.get(f"{self.endpoint}{own.id}/imagen/").status_code, 200)
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.get(self.endpoint).data["count"], 2)
        self.assertEqual(self.client.get(f"{self.endpoint}{other.id}/imagen/").status_code, 200)

    def test_filters_individually_and_combined_with_chilean_dates(self):
        first = self.expense()
        second = self.expense(self.other, instalacion="Planta Sur", persona="María Díaz")
        # In September Chile is UTC-3: 02:59 UTC belongs to the previous local day.
        RendicionGasto.objects.filter(pk=first.pk).update(creada_en=datetime(2026, 9, 25, 2, 59, tzinfo=timezone.utc))
        RendicionGasto.objects.filter(pk=second.pk).update(creada_en=datetime(2026, 9, 25, 3, 0, tzinfo=timezone.utc))
        self.client.force_authenticate(self.admin)
        for filters in [
            {"supervisor": self.supervisor.id}, {"instalacion": "norte"}, {"persona": "Pedro"},
            {"fecha_desde": "2026-09-24", "fecha_hasta": "2026-09-24"},
            {"supervisor": self.supervisor.id, "instalacion": "Norte", "persona": "Soto", "fecha_desde": "2026-09-24", "fecha_hasta": "2026-09-24"},
        ]:
            with self.subTest(filters=filters):
                response = self.client.get(self.endpoint, filters)
                self.assertEqual(response.status_code, 200)
                self.assertEqual([item["id"] for item in response.data["results"]], [first.id])
        self.assertEqual(self.client.get(self.endpoint, {"instalacion": "Norte", "persona": "María"}).data["count"], 0)

    def test_invalid_filters_return_validation_error(self):
        for filters in [{"supervisor": "invalid"}, {"fecha_desde": "invalid"}, {"fecha_desde": "2026-09-25", "fecha_hasta": "2026-09-24"}]:
            self.assertEqual(self.client.get(self.endpoint, filters).status_code, 400)

    def test_filter_options_are_scoped(self):
        self.expense()
        self.expense(self.other)
        response = self.client.get(f"{self.endpoint}filtros/")
        self.assertEqual(len(response.data["supervisores"]), 1)
        self.client.force_authenticate(self.admin)
        self.assertEqual(len(self.client.get(f"{self.endpoint}filtros/").data["supervisores"]), 2)

    def test_authentication_required_and_records_cannot_be_modified(self):
        expense = self.expense()
        self.assertEqual(self.client.patch(f"{self.endpoint}{expense.id}/", {"persona": "Alterado"}).status_code, 405)
        self.assertEqual(self.client.delete(f"{self.endpoint}{expense.id}/").status_code, 405)
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get(self.endpoint).status_code, 401)
        self.assertEqual(self.client.post(self.endpoint, {}, format="multipart").status_code, 401)

    def test_pagination(self):
        for _ in range(21):
            self.expense()
        response = self.client.get(self.endpoint)
        self.assertEqual(response.data["count"], 21)
        self.assertEqual(len(response.data["results"]), 20)
        self.assertEqual(len(self.client.get(self.endpoint, {"page": 2}).data["results"]), 1)

    @patch("supervision.rendiciones.upload_document")
    def test_amount_and_reason_are_required_and_validated_before_upload(self, upload):
        cases = [{"monto": None}, {"motivo": None}, {"monto": "0"}, {"monto": "-1"},
                 {"monto": "1.5"}, {"monto": "abc"}, {"monto": "1000000000000"},
                 {"motivo": "  "}, {"motivo": "x" * 1001}]
        for changes in cases:
            payload = {"imagen": self.image(), "instalacion": "Planta", "persona": "Pedro", "monto": "15000", "motivo": "Traslado"}
            payload.update(changes)
            payload = {key: value for key, value in payload.items() if value is not None}
            with self.subTest(changes=changes):
                response = self.client.post(self.endpoint, payload, format="multipart")
                self.assertEqual(response.status_code, 400, response.data)
        upload.assert_not_called()

    def report(self, formato="xlsx", **filters):
        params = {"fecha_desde": "2000-01-01", "fecha_hasta": "2100-12-31", "formato": formato, **filters}
        return self.client.get(f"{self.endpoint}informe/", params)

    def test_totals_include_all_pages_and_mark_legacy_records(self):
        for _ in range(21):
            self.expense(monto=1234)
        self.expense(monto=None, motivo="")
        self.expense(self.other, monto=900000)
        response = self.client.get(self.endpoint)
        self.assertEqual(response.data["resumen"], {"total": "25914", "moneda": "CLP", "con_monto": 21, "sin_monto": 1})
        self.assertEqual(len(response.data["results"]), 20)
        book = load_workbook(BytesIO(self.report().content))
        self.assertEqual(book.sheetnames, ["Resumen", "Detalle", "Por supervisor", "Por instalación", "Por día"])
        self.assertEqual(book["Detalle"].max_row, 24)  # Header, 22 expenses, total.
        self.assertEqual(book["Detalle"].cell(24, 7).value, 25914)
        self.assertEqual(book["Por instalación"]["C2"].value, 25914)
        self.assertEqual(book["Por instalación"]["D2"].value, 1)
        self.assertEqual(book["Detalle"].freeze_panes, "A2")

    def test_excel_preserves_text_numeric_amounts_and_local_dates(self):
        row = self.expense(monto=999999999999, motivo="=HYPERLINK(\"https://example.test\")")
        RendicionGasto.objects.filter(pk=row.pk).update(creada_en=datetime(2026, 9, 25, 2, 59, tzinfo=timezone.utc))
        response = self.report(fecha_desde="2026-09-24", fecha_hasta="2026-09-24")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book["Detalle"]["B2"].value, datetime(2026, 9, 24, 23, 59))
        self.assertEqual(book["Detalle"]["F2"].data_type, "s")
        self.assertEqual(book["Detalle"]["F2"].value, row.motivo)
        self.assertEqual(book["Detalle"]["G2"].data_type, "n")
        self.assertEqual(book["Detalle"]["G2"].value, 999999999999)

    def test_reports_respect_filters_scope_and_chilean_day_boundary(self):
        first = self.expense()
        second = self.expense(self.other, instalacion="Planta Sur")
        RendicionGasto.objects.filter(pk=first.pk).update(creada_en=datetime(2026, 9, 25, 2, 59, tzinfo=timezone.utc))
        RendicionGasto.objects.filter(pk=second.pk).update(creada_en=datetime(2026, 9, 25, 3, 0, tzinfo=timezone.utc))
        self.assertEqual(load_workbook(BytesIO(self.report(supervisor=self.other.id).content))["Detalle"].max_row, 2)
        self.client.force_authenticate(self.admin)
        for filters in [{"fecha_desde": "2026-09-24", "fecha_hasta": "2026-09-24"},
                        {"supervisor": self.supervisor.id, "instalacion": "Norte", "persona": "Pedro"}]:
            book = load_workbook(BytesIO(self.report(**filters).content))
            self.assertEqual(book["Detalle"].max_row, 3)
            self.assertEqual(book["Detalle"]["A2"].value, first.id)

    def test_pdf_handles_multiple_pages_long_reasons_and_totals(self):
        for _ in range(25):
            self.expense(monto=2000, motivo="Traslado <guardia> & alimentación. " * 25)
        response = self.report("pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        document = PdfReader(BytesIO(response.content))
        self.assertGreater(len(document.pages), 1)
        text = " ".join(page.extract_text() for page in document.pages)
        self.assertIn("50.000", text)
        self.assertIn("Por supervisor", text)
        self.assertIn("Detalle de rendiciones", text)
        self.assertIn("Traslado <guardia> &", text)

    def test_empty_reports_and_invalid_report_requests(self):
        for formato in ["pdf", "xlsx"]:
            self.assertEqual(self.report(formato).status_code, 200)
        for params in [{}, {"fecha_desde": "2026-01-01"},
                       {"fecha_desde": "2026-02-01", "fecha_hasta": "2026-01-01"},
                       {"fecha_desde": "invalid", "fecha_hasta": "2026-01-01"}]:
            self.assertEqual(self.client.get(f"{self.endpoint}informe/", params).status_code, 400)
        self.assertEqual(self.report("csv").status_code, 400)
        self.client.force_authenticate(None)
        self.assertEqual(self.report().status_code, 401)
