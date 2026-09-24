"""Informes completos (sin paginación) con importes numéricos en CLP."""
from collections import defaultdict
from io import BytesIO
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

CHILE = ZoneInfo("America/Santiago")
MONEY_FORMAT = '"$" #,##0;[Red]-"$" #,##0'


def money(value):
    return "Sin monto" if value is None else "$ " + f"{value:,}".replace(",", ".")


def expense_report(queryset, filters, formato):
    rows = list(queryset)
    total = sum(int(row.monto) for row in rows if row.monto is not None)
    missing = sum(row.monto is None for row in rows)
    generated = timezone.localtime(timezone.now(), CHILE).strftime("%d/%m/%Y %H:%M")
    metadata = [
        f"Período: {filters['fecha_desde']} al {filters['fecha_hasta']} (ambas fechas incluidas)",
        "Fecha de referencia: guardado de la rendición · Hora de Chile · Moneda: CLP",
        f"Emitido: {generated} · Registros: {len(rows)} · Sin monto: {missing}",
    ]
    for key, label in [("supervisor", "Supervisor ID"), ("instalacion", "Instalación contiene"), ("persona", "Persona contiene")]:
        if filters.get(key):
            metadata.append(f"{label}: {filters[key]}")
    metadata.append("Los registros sin monto no se incluyen en los totales. El total corresponde al monto registrado, sin desglose de impuestos.")

    groups = {"Por supervisor": defaultdict(lambda: [0, 0, 0]),
              "Por instalación": defaultdict(lambda: [0, 0, 0]),
              "Por día": defaultdict(lambda: [0, 0, 0])}
    details = []
    for row in rows:
        local = timezone.localtime(row.creada_en, CHILE)
        amount = int(row.monto) if row.monto is not None else None
        details.append([row.id, local.replace(tzinfo=None), row.supervisor_nombre, row.instalacion,
                        row.persona, row.motivo or "Sin motivo registrado", amount])
        keys = [f"{row.supervisor_nombre} (#{row.supervisor_id})", row.instalacion, local.strftime("%Y-%m-%d")]
        for group, key in zip(groups.values(), keys):
            group[key][0] += 1
            group[key][1] += amount or 0
            group[key][2] += amount is None
    summaries = {title: [[key, *values] for key, values in sorted(group.items())]
                 for title, group in groups.items()}
    content = (excel_report(metadata, details, summaries, total, missing) if formato == "xlsx"
               else pdf_report(metadata, details, summaries, total))
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if formato == "xlsx" else "application/pdf"
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="rendiciones-{filters["fecha_desde"]}-{filters["fecha_hasta"]}.{formato}"'
    response["Cache-Control"] = "private, no-store"
    return response


def excel_report(metadata, details, summaries, total, missing):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary.append(["INOUT · Informe de rendiciones"])
    for line in metadata:
        summary.append([line])
    summary.append([])
    summary.append(["Registros", len(details)])
    summary.append(["Con monto", len(details) - missing])
    summary.append(["Sin monto", missing])
    summary.append(["TOTAL CLP", total])
    summary.cell(summary.max_row, 2).number_format = MONEY_FORMAT
    summary.column_dimensions["A"].width = 95
    summary.column_dimensions["B"].width = 22
    summary.sheet_properties.pageSetUpPr.fitToPage = True

    def add_sheet(name, headers, rows, widths, amount_column):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"
        sheet.append(["TOTAL CLP"] + [None] * (amount_column - 2) + [total])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        for row in sheet.iter_rows(min_row=2):
            row[amount_column - 1].number_format = MONEY_FORMAT
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        return sheet

    detail = add_sheet("Detalle", ["ID", "Fecha y hora (Chile)", "Supervisor", "Instalación", "Persona", "Motivo", "Monto CLP"],
                       details, [12, 23, 30, 32, 30, 65, 22], 7)
    for row in detail.iter_rows(min_row=2, max_row=len(details) + 1):
        row[1].number_format = "dd/mm/yyyy hh:mm"
    for title, rows in summaries.items():
        sheet = add_sheet(title, [title, "Cantidad", "Total CLP", "Sin monto"], rows, [55, 16, 22, 16], 3)
        sheet.cell(sheet.max_row, 2, len(details))
        sheet.cell(sheet.max_row, 4, missing)
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                # User text must remain text, even if it begins with '='.
                if isinstance(cell.value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="163451")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf_report(metadata, details, summaries, total):
    output = BytesIO()
    styles = getSampleStyleSheet()
    styles["BodyText"].fontSize = 8
    styles["BodyText"].leading = 11
    story = [Paragraph("INOUT · Informe de rendiciones", styles["Title"])]

    def paragraph(value):
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), styles["BodyText"])

    story.extend(paragraph(line) for line in metadata)
    story.extend([Spacer(1, 12), Paragraph(f"Total registrado: {money(total)} CLP", styles["Heading2"])])

    def table(headers, rows, widths):
        data = [[paragraph(value) for value in headers]]
        data.extend([[paragraph(value) for value in row] for row in rows])
        result = LongTable(data, colWidths=widths, repeatRows=1, splitInRow=1, hAlign="LEFT")
        result.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dce8f2")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f6f9")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("LINEBELOW", (0, 0), (-1, 0), .5, colors.grey),
        ]))
        return result

    for title, rows in summaries.items():
        story.append(Paragraph(title, styles["Heading2"]))
        data = [[key, count, money(amount), absent] for key, count, amount, absent in rows]
        story.append(table([title, "Cantidad", "Total CLP", "Sin monto"], data, [440, 90, 140, 100]))
    story.append(Paragraph("Detalle de rendiciones", styles["Heading2"]))
    if details:
        data = [[id_, date.strftime("%d/%m/%Y %H:%M"), sup, site, person, reason, money(amount)]
                for id_, date, sup, site, person, reason, amount in details]
        data.append(["TOTAL", "", "", "", "", "", money(total)])
        story.append(table(["ID", "Guardado (Chile)", "Supervisor", "Instalación", "Persona", "Motivo", "Monto CLP"],
                           data, [40, 80, 105, 105, 100, 240, 100]))
    else:
        story.append(paragraph("No hay rendiciones para el período y filtros seleccionados."))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(36, 20, "INOUT · Rendiciones de gastos · CLP")
        canvas.drawRightString(806, 20, f"Página {document.page}")
        canvas.restoreState()

    SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=36, leftMargin=36,
                      topMargin=30, bottomMargin=36, title="Informe de rendiciones INOUT").build(
                          story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
